"""
Chart Engine — creates openpyxl chart objects based on ChartSpec and data ranges.
"""

from openpyxl.chart import (
    BarChart, LineChart, PieChart, DoughnutChart, AreaChart, Reference,
    BarChart3D,
)
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.utils import get_column_letter
from schemas.schemas import ChartSpec


# Chart type mapping with fallback rules
CHART_TYPE_MAP = {
    "column": "column",
    "bar": "bar",
    "line": "line",
    "pie": "pie",
    "doughnut": "doughnut",
    "area": "area",
    "combo": "column",  # Combo falls back to column in openpyxl
}

# Default chart type recommendations based on analysis pattern
ANALYSIS_CHART_DEFAULTS = {
    "by_month": "line",
    "by_region": "column",
    "by_product": "bar",
    "by_salesperson": "column",
    "by_department": "bar",
    "by_category": "bar",
    "market_share": "doughnut",
    "trend": "line",
    "comparison": "column",
    "distribution": "bar",
}


class ChartEngine:
    """Creates real Excel chart objects using openpyxl."""

    @staticmethod
    def recommend_chart_type(analysis_name: str, group_by: list[str]) -> str:
        """Recommend a chart type based on the analysis pattern."""
        name_lower = analysis_name.lower()

        for pattern, chart_type in ANALYSIS_CHART_DEFAULTS.items():
            if pattern.replace("_", " ") in name_lower or pattern.replace("_", " ") in " ".join(g.lower() for g in group_by):
                return chart_type

        # Default based on group_by keywords
        if group_by:
            group_lower = group_by[0].lower()
            if "month" in group_lower or "date" in group_lower or "year" in group_lower:
                return "line"
            if "region" in group_lower or "area" in group_lower:
                return "column"

        return "column"

    @staticmethod
    def create_chart(
        spec: ChartSpec,
        ws,
        data_start_row: int,
        data_end_row: int,
        category_col: int,
        value_col: int,
        position: str = "A1",
    ):
        """
        Create an openpyxl chart object and add it to the worksheet.

        Args:
            spec: Chart specification
            ws: openpyxl worksheet
            data_start_row: First row of data (1-indexed, including header)
            data_end_row: Last row of data
            category_col: Column number for categories (1-indexed)
            value_col: Column number for values (1-indexed)
            position: Cell position for the chart (e.g., "A1")
        """
        chart_type = CHART_TYPE_MAP.get(spec.chart_type, "column")

        if chart_type == "column":
            chart = BarChart()
            chart.type = "col"
            chart.grouping = "clustered"
        elif chart_type == "bar":
            chart = BarChart()
            chart.type = "bar"
            chart.grouping = "clustered"
        elif chart_type == "line":
            chart = LineChart()
            chart.grouping = "standard"
        elif chart_type == "pie":
            chart = PieChart()
        elif chart_type == "doughnut":
            chart = DoughnutChart()
        elif chart_type == "area":
            chart = AreaChart()
            chart.grouping = "standard"
        else:
            chart = BarChart()
            chart.type = "col"
            chart.grouping = "clustered"

        # Set chart properties
        chart.title = spec.title
        chart.width = spec.width
        chart.height = spec.height

        # Set axis titles for applicable chart types
        if chart_type not in ("pie", "doughnut"):
            if spec.x_axis_title:
                chart.x_axis.title = spec.x_axis_title
            if spec.y_axis_title:
                chart.y_axis.title = spec.y_axis_title

        # Chart styling
        chart.style = 10  # A clean, professional style

        # Data reference
        data_ref = Reference(
            ws,
            min_col=value_col,
            min_row=data_start_row,
            max_row=data_end_row,
        )

        # Category reference
        cats_ref = Reference(
            ws,
            min_col=category_col,
            min_row=data_start_row + 1,  # Skip header
            max_row=data_end_row,
        )

        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)

        # Add data labels for pie/doughnut
        if chart_type in ("pie", "doughnut"):
            chart.dataLabels = DataLabelList()
            chart.dataLabels.showPercent = True
            chart.dataLabels.showCatName = True

        ws.add_chart(chart, position)
        return chart

    @staticmethod
    def create_multi_series_chart(
        spec: ChartSpec,
        ws,
        data_start_row: int,
        data_end_row: int,
        category_col: int,
        value_start_col: int,
        value_end_col: int,
        position: str = "A1",
    ):
        """Create a chart with multiple data series (for cross-tab data)."""
        chart_type = CHART_TYPE_MAP.get(spec.chart_type, "column")

        if chart_type == "line":
            chart = LineChart()
            chart.grouping = "standard"
        elif chart_type == "bar":
            chart = BarChart()
            chart.type = "bar"
            chart.grouping = "clustered"
        else:
            chart = BarChart()
            chart.type = "col"
            chart.grouping = "clustered"

        chart.title = spec.title
        chart.width = spec.width
        chart.height = spec.height
        chart.style = 10

        if spec.x_axis_title:
            chart.x_axis.title = spec.x_axis_title
        if spec.y_axis_title:
            chart.y_axis.title = spec.y_axis_title

        # Add data for each series
        data_ref = Reference(
            ws,
            min_col=value_start_col,
            max_col=value_end_col,
            min_row=data_start_row,
            max_row=data_end_row,
        )
        cats_ref = Reference(
            ws,
            min_col=category_col,
            min_row=data_start_row + 1,
            max_row=data_end_row,
        )

        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)

        ws.add_chart(chart, position)
        return chart


chart_engine = ChartEngine()
