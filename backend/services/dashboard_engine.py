"""
Dashboard Engine — dynamically lays out KPIs, charts, and filters
on a dashboard sheet using a grid-based system.
"""

from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from schemas.schemas import DashboardSpec, KPISpec
import pandas as pd
import numpy as np


# Dashboard color palette
COLORS = {
    "primary": "1B2A4A",       # Dark navy
    "secondary": "2E86AB",     # Teal blue
    "accent1": "A23B72",       # Berry
    "accent2": "F18F01",       # Orange
    "accent3": "2CA58D",       # Green
    "accent4": "5C4D7D",       # Purple
    "kpi_bg": "F0F4F8",        # Light gray-blue
    "kpi_border": "D1D9E6",    # Border gray
    "white": "FFFFFF",
    "dark_text": "1B2A4A",
    "light_text": "6B7B8D",
    "title_bg": "1B2A4A",
    "title_text": "FFFFFF",
}

KPI_COLORS = ["2E86AB", "A23B72", "F18F01", "2CA58D", "5C4D7D", "E84855"]


class DashboardEngine:
    """Generates a dynamic dashboard layout on an Excel worksheet."""

    def __init__(self, ws):
        self.ws = ws
        self.current_row = 1
        self.max_col = 14  # Columns A through N

    def _set_column_widths(self):
        """Set consistent column widths for the dashboard grid."""
        for i in range(1, self.max_col + 1):
            self.ws.column_dimensions[get_column_letter(i)].width = 12

    def _merge_and_style(
        self,
        start_row: int,
        start_col: int,
        end_row: int,
        end_col: int,
        value: str,
        font: Font = None,
        alignment: Alignment = None,
        fill: PatternFill = None,
        border: Border = None,
    ):
        """Merge cells and apply styling."""
        start_cell = f"{get_column_letter(start_col)}{start_row}"
        end_cell = f"{get_column_letter(end_col)}{end_row}"
        self.ws.merge_cells(f"{start_cell}:{end_cell}")

        cell = self.ws[start_cell]
        cell.value = value
        if font:
            cell.font = font
        if alignment:
            cell.alignment = alignment
        if fill:
            cell.fill = fill
        if border:
            cell.border = border

    def write_title(self, title: str):
        """Write the dashboard title bar."""
        self._merge_and_style(
            self.current_row, 1, self.current_row + 1, self.max_col,
            value=title.upper(),
            font=Font(name="Calibri", size=18, bold=True, color=COLORS["title_text"]),
            alignment=Alignment(horizontal="center", vertical="center"),
            fill=PatternFill(start_color=COLORS["title_bg"], end_color=COLORS["title_bg"], fill_type="solid"),
        )
        self.ws.row_dimensions[self.current_row].height = 25
        self.ws.row_dimensions[self.current_row + 1].height = 25
        self.current_row += 3  # Title + spacing

    def write_kpis(self, kpis: list[dict]):
        """
        Write KPI cards across the dashboard.
        kpis: list of {"label": str, "value": any, "format": str}
        """
        if not kpis:
            return

        num_kpis = len(kpis)
        cols_per_kpi = max(2, self.max_col // num_kpis)
        thin_border = Border(
            left=Side(style="thin", color=COLORS["kpi_border"]),
            right=Side(style="thin", color=COLORS["kpi_border"]),
            top=Side(style="thin", color=COLORS["kpi_border"]),
            bottom=Side(style="thin", color=COLORS["kpi_border"]),
        )

        for i, kpi in enumerate(kpis):
            start_col = 1 + i * cols_per_kpi
            end_col = min(start_col + cols_per_kpi - 1, self.max_col)
            color = KPI_COLORS[i % len(KPI_COLORS)]

            # KPI Label
            self._merge_and_style(
                self.current_row, start_col, self.current_row, end_col,
                value=kpi["label"].upper(),
                font=Font(name="Calibri", size=9, bold=True, color=COLORS["light_text"]),
                alignment=Alignment(horizontal="center", vertical="center"),
                fill=PatternFill(start_color=COLORS["kpi_bg"], end_color=COLORS["kpi_bg"], fill_type="solid"),
                border=thin_border,
            )

            # KPI Value & Dynamic Formula
            formula = kpi.get("formula")
            value = kpi.get("value")
            fmt = kpi.get("format", "number")

            num_formats = {
                "currency": "$#,##0",
                "currency_decimal": "$#,##0.00",
                "percentage": "0.0%",
                "integer": "#,##0",
                "number": "#,##0",
                "number_decimal": "#,##0.00",
            }
            excel_num_format = num_formats.get(fmt, "#,##0")

            cell_val = formula if formula else (
                f"${value:,.0f}" if fmt == "currency" and isinstance(value, (int, float))
                else (f"{value:.1%}" if fmt == "percentage" and isinstance(value, (int, float))
                else (f"{value:,.0f}" if isinstance(value, (int, float)) else str(value)))
            )

            self._merge_and_style(
                self.current_row + 1, start_col, self.current_row + 2, end_col,
                value=cell_val,
                font=Font(name="Calibri", size=22, bold=True, color=color),
                alignment=Alignment(horizontal="center", vertical="center"),
                fill=PatternFill(start_color=COLORS["white"], end_color=COLORS["white"], fill_type="solid"),
                border=thin_border,
            )

            start_cell_ref = f"{get_column_letter(start_col)}{self.current_row + 1}"
            if formula:
                self.ws[start_cell_ref].number_format = excel_num_format

        self.ws.row_dimensions[self.current_row].height = 20
        self.ws.row_dimensions[self.current_row + 1].height = 30
        self.ws.row_dimensions[self.current_row + 2].height = 15
        self.current_row += 4  # KPIs + spacing

    def get_chart_position(self, chart_index: int, charts_per_row: int = 2) -> str:
        """Calculate chart position based on index in the chart grid."""
        row_in_grid = chart_index // charts_per_row
        col_in_grid = chart_index % charts_per_row

        chart_height_rows = 15
        chart_row = self.current_row + row_in_grid * (chart_height_rows + 1)

        if charts_per_row == 1:
            chart_col = 1
        else:
            cols_per_chart = self.max_col // charts_per_row
            chart_col = 1 + col_in_grid * cols_per_chart

        return f"{get_column_letter(chart_col)}{chart_row}"

    def advance_past_charts(self, num_charts: int, charts_per_row: int = 2):
        """Advance current_row past the chart grid area."""
        chart_height_rows = 15
        num_rows_of_charts = (num_charts + charts_per_row - 1) // charts_per_row
        self.current_row += num_rows_of_charts * (chart_height_rows + 1) + 1

    def write_section_header(self, title: str):
        """Write a section header."""
        self._merge_and_style(
            self.current_row, 1, self.current_row, self.max_col,
            value=title,
            font=Font(name="Calibri", size=12, bold=True, color=COLORS["primary"]),
            alignment=Alignment(horizontal="left", vertical="center"),
        )
        self.current_row += 1


dashboard_engine_cls = DashboardEngine
